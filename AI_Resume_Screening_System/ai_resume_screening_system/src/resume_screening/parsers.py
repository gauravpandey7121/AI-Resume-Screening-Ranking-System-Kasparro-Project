from __future__ import annotations

import hashlib
import re
from pathlib import Path

from pydantic import BaseModel

from .models import Project, ResumeData

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

SECTION_ALIASES = {
    "summary": {"summary", "professional summary", "profile", "objective", "career objective"},
    "skills": {"skills", "technical skills", "key skills", "technical skill", "skills & technologies"},
    "projects": {"projects", "key projects", "project experience", "projects & experience"},
    "experience": {"experience", "work experience", "professional experience", "employment", "internship", "internships"},
    "education": {"education", "academic background"},
    "certifications": {"certifications", "certificates"},
}

HEADING_RE = re.compile(r"^[A-Za-z][A-Za-z0-9 /&+().-]{1,60}$")
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d[\d ()-]{8,}\d)(?!\d)")
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9-]+)", re.I)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _extract_pdf(path: Path) -> str:
    errors: list[str] = []
    try:
        import fitz  # PyMuPDF

        with fitz.open(path) as doc:
            text = "\n".join(page.get_text("text") for page in doc)
        if text.strip():
            return text
    except Exception as exc:
        errors.append(f"PyMuPDF: {exc}")

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path), strict=False)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return text
    except Exception as exc:
        errors.append(f"pypdf: {exc}")

    details = "; ".join(errors) if errors else "no text extracted"
    raise ValueError(f"Unable to extract PDF text: {details}")


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)


def _extract_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        return _extract_pdf(path)
    if path.suffix.lower() == ".docx":
        return _extract_docx(path)
    if path.suffix.lower() == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {path.suffix}")


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _normalized_heading(line: str) -> str | None:
    clean = re.sub(r"[^A-Za-z0-9 /&+().-]", "", line).strip().lower()
    if clean in SECTION_ALIASES["summary"]:
        return "summary"
    for canonical in ("skills", "projects", "experience", "education", "certifications"):
        if clean in SECTION_ALIASES[canonical]:
            return canonical
    return None


def extract_sections(text: str) -> dict[str, str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in lines:
        heading = _normalized_heading(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        if current:
            sections[current].append(line)
    return {key: "\n".join(values) for key, values in sections.items()}


def _clean_candidate_name(text: str, email: str | None) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    bad = {
        "summary", "professional summary", "skills", "technical skills", "experience",
        "professional experience", "education", "projects", "objective", "get in touch",
        "contact", "contact me", "core competencies", "impact highlights", "software engineer", "core", "competencies", "impact", "highlights", "get in touch", "cgpa", "job", "simulation", "higher", "secondary", "pcmb",
        "full stack developer", "backend developer", "ai engineer", "data analyst", "developer", "software", "engineering",
    }

    def clean(line: str) -> str:
        return re.sub(r"\s+", " ", line).strip(" |•-!,:;")

    def usable_name(candidate: str) -> bool:
        tokens = candidate.split()
        if not (1 <= len(tokens) <= 5 and 4 <= len(candidate) <= 50):
            return False
        lower_tokens = {re.sub(r"[^a-z]", "", t.casefold()) for t in tokens}
        if EMAIL_RE.search(candidate) or "http" in candidate.lower() or candidate.casefold() in bad:
            return False
        if lower_tokens & bad:
            return False
        if _normalized_heading(candidate):
            return False
        lower = candidate.casefold()
        if any(role in lower for role in ("software engineer", "developer", "analyst", "student", "graduate")):
            return False
        letters = [ch for ch in candidate if ch.isalpha()]
        return len(letters) >= 4

    # First preference: a natural name at the top of the resume/contact block.
    for line in lines[:6]:
        candidate = clean(line)
        tokens = candidate.split()
        title_case_like = all((t[:1].isupper() or re.match(r"^[A-Z]\.?$", t)) for t in tokens if t)
        uppercase_like = candidate.upper() == candidate
        if (usable_name(candidate) and len(tokens) <= 4 and (title_case_like or uppercase_like)):
            return candidate

    email_idx = next((i for i, line in enumerate(lines) if EMAIL_RE.search(line)), -1)
    search_ranges = []
    if email_idx >= 0:
        search_ranges.append(range(max(0, email_idx - 10), min(len(lines), email_idx + 11)))
    search_ranges.append(range(0, min(20, len(lines))))
    for indices in search_ranges:
        for i in indices:
            candidate = clean(lines[i])
            letters = [ch for ch in candidate if ch.isalpha()]
            upper_ratio = sum(ch.isupper() for ch in letters) / max(len(letters), 1)
            if usable_name(candidate) and upper_ratio >= 0.75:
                return candidate

    if email:
        local = email.split("@", 1)[0]
        local = re.sub(r"\d+$", "", local).replace(".", " ").replace("_", " ").replace("-", " ")
        guess = " ".join(part.capitalize() for part in local.split())
        if guess:
            return guess
    return "Unknown Candidate"


def _split_skill_items(text: str) -> list[str]:
    chunks = re.split(r"\n|•|,|;|\|", text)
    items: list[str] = []
    for chunk in chunks:
        item = re.sub(r"^[^:]{0,40}:\s*", "", chunk.strip())
        item = re.sub(r"\s+", " ", item).strip(" -–—")
        if 1 < len(item) < 50 and len(item.split()) <= 8:
            items.append(item)
    return _dedupe(items)


def _dedupe(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = value.casefold().strip()
        if key and key not in seen:
            seen.add(key)
            result.append(value.strip())
    return result


def _extract_projects(section: str) -> list[Project]:
    if not section:
        return []
    lines = [re.sub(r"^[-•●▪\s]+", "", line).strip() for line in section.splitlines() if line.strip()]
    projects: list[Project] = []
    current: Project | None = None
    tech_re = re.compile(r"(?:[|–-]|\bTech(?:nologies| Stack)?\s*:)\s*(.+)$", re.I)
    for line in lines:
        looks_like_header = (
            len(line) <= 120
            and not line.lower().startswith(("built ", "developed ", "implemented ", "designed ", "created ", "worked ", "deployed ", "achieved ", "integrated "))
            and ("|" in line or ":" in line or (len(line.split()) <= 12 and not line.endswith(".")))
        )
        if looks_like_header:
            if current:
                projects.append(current)
            techs: list[str] = []
            match = tech_re.search(line)
            if match:
                techs = _split_skill_items(match.group(1))
                name = line[:match.start()].strip(" |:-")
            else:
                name = line.split("|")[0].strip()
            current = Project(name=name or "Project", technologies=techs, evidence=[])
            continue
        if current is None:
            current = Project(name="Project", evidence=[])
        current.evidence.append(line)
        tech_match = tech_re.search(line)
        if tech_match:
            current.technologies.extend(_split_skill_items(tech_match.group(1)))
    if current:
        projects.append(current)

    for project in projects:
        project.technologies = _dedupe(project.technologies)
        project.description = " ".join(project.evidence[:3]).strip()
    return [p for p in projects if p.name or p.evidence]


def parse_resume(path: Path) -> ResumeData:
    raw = _normalize(_extract_text(path))
    if not raw:
        raise ValueError("Extracted resume text is empty")
    sections = extract_sections(raw)
    email = (EMAIL_RE.search(raw).group(0) if EMAIL_RE.search(raw) else None)
    phone = (PHONE_RE.search(raw).group(0) if PHONE_RE.search(raw) else None)
    github_match = GITHUB_RE.search(raw)
    github_url = f"https://github.com/{github_match.group(1)}" if github_match else None
    skills = _split_skill_items(sections.get("skills", ""))
    projects = _extract_projects(sections.get("projects", ""))
    experience_lines = [re.sub(r"\s+", " ", x).strip(" -•●") for x in sections.get("experience", "").splitlines() if x.strip()]
    education_lines = [re.sub(r"\s+", " ", x).strip(" -•●") for x in sections.get("education", "").splitlines() if x.strip()]
    return ResumeData(
        file_name=path.name,
        file_hash=file_sha256(path),
        candidate_name=_clean_candidate_name(raw, email),
        email=email,
        phone=phone,
        github_url=github_url,
        skills=skills,
        projects=projects,
        experience=experience_lines,
        education=education_lines,
        sections=sections,
        raw_text=raw,
    )
